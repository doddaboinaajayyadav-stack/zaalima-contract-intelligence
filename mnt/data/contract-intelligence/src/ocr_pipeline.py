from pathlib import Path
from typing import BinaryIO


def extract_text_from_pdf(path: str | Path) -> str:
    import pdfplumber

    path = Path(path)
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
            else:
                # OCR fallback for scanned/image-only pages.
                try:
                    import pytesseract
                    image = page.to_image(resolution=200).original
                    pages.append(pytesseract.image_to_string(image))
                except Exception as exc:
                    pages.append(f"[OCR unavailable for page: {exc}]")
    return "\n\n".join(pages).strip()


def extract_text_from_docx(path: str | Path) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts).strip()


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".docx":
        return extract_text_from_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {suffix}. Use PDF, DOCX, TXT or MD.")
